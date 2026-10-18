# -*- coding: utf-8 -*-
"""PPT individual por supervisor (5 laminas)."""
import sys
import pandas as pd
from pathlib import Path
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from style import *
from supervisor_data import bundle, SUPERVISORES
import charts_supervisor as ch

SC = Path(__file__).parent
IMG = SC / "img"
DESTDIR = Path(r"C:\Users\Rudy\OneDrive - AM COMERCIAL\Rudy\Septiembre 2026\Seguimiento")

def build(key):
    b = bundle(key)
    cfg = b["cfg"]; nombre = cfg["nombre"]
    ch.presu_por_vendedor(key, b["por_vend"])
    ch.novis_por_vendedor(key, b["por_vend_vis"])

    prs = nueva_presentacion()
    BLANK = prs.slide_layouts[6]
    PIE = f"AM Comercial \u00b7 Confidencial \u00b7 Reuni\u00f3n con {nombre} ({cfg['detalle']}) \u00b7 Corte 19-sep-2026"

    # SLIDE 1 - Portada
    sl = prs.slides.add_slide(BLANK)
    rect(sl, 0, 0, 13.333, 7.5, fill=AZUL)
    rect(sl, 0, 4.55, 13.333, 0.09, fill=VERDE)
    if LOGO.exists(): sl.shapes.add_picture(str(LOGO), Inches(0.6), Inches(0.55), height=Inches(0.62))
    tb(sl, 0.6, 2.35, 12.1, 1.0, f"Reuni\u00f3n con {nombre}", size=34, bold=True, color=BLANCO)
    tb(sl, 0.6, 3.25, 12.1, 0.5, cfg["detalle"], size=15, color=RGBColor(0xD8,0xE0,0xF5))
    tb(sl, 0.6, 4.85, 12.1, 0.4, "Corte al 19 de septiembre 2026 (d\u00eda 19 de 30)  \u00b7  Fuente: MovilBusiness5STD", size=11.5, color=RGBColor(0xC7,0xD1,0xEE))
    tb(sl, 0.6, 6.85, 12.1, 0.3, "AM Comercial \u00b7 Confidencial", size=9, color=RGBColor(0xAE,0xBA,0xE0))

    # SLIDE 2 - Presupuesto por vendedor
    sl = prs.slides.add_slide(BLANK)
    head(sl, "1. Presupuesto vs. venta al d\u00eda por vendedor", "Presupuesto al d\u00eda prorrateado seg\u00fan d\u00edas h\u00e1biles transcurridos del mes.", 2, PIE)
    t = b["tot"]
    kpi(sl, 0.5, 1.35, 2.95, money(t.PresupuestoAlDia/1e6)+"MM", "Presupuesto al d\u00eda (RD$)", AZUL)
    kpi(sl, 3.62, 1.35, 2.95, money(t.Ejecutado/1e6)+"MM", "Ejecutado a la fecha", VERDE)
    color_alc = ROJO if t.PctAlcance<80 else (AMBAR if t.PctAlcance<95 else VERDE)
    kpi(sl, 6.74, 1.35, 2.95, pct(t.PctAlcance), "Alcance al d\u00eda", color_alc)
    kpi(sl, 9.86, 1.35, 2.99, money(t.Proyectado/1e6)+"MM", "Proyecci\u00f3n de cierre", GRIS2)
    sl.shapes.add_picture(str(IMG/f"{key}_presu_vend.png"), Inches(0.5), Inches(2.55), width=Inches(12.35))
    col = b["por_suplidor"][b["por_suplidor"].Suplidor.str.contains("COLOMBINA", case=False, na=False)]
    if len(col):
        c = col.iloc[0]
        lectura(sl, f"Colombina: RD$ {money(c.Ejecutado)} ejecutado de RD$ {money(c.PresupuestoAlDia)} al d\u00eda ({pct(c.PctAlcance)} de alcance).",
                color=RCLARO if c.PctAlcance<80 else (ACLARO if c.PctAlcance<95 else GCLARO), barra=(ROJO if c.PctAlcance<80 else (AMBAR if c.PctAlcance<95 else VERDE)))

    # SLIDE 3 - Capacidad de compra
    sl = prs.slides.add_slide(BLANK)
    head(sl, "2. Clientes de mayor capacidad de compra vs. septiembre", "Top clientes por promedio mensual (marzo-agosto 2026) vs. lo comprado en septiembre a la fecha.", 3, PIE)
    top = b["topN"].head(12)
    rows=[["Vendedor","Cliente","Promedio\nmensual","Esperado\na la fecha","Comprado\nSept.","%\nCumpl."]]
    cc={}
    for i,r in enumerate(top.itertuples(), start=1):
        rows.append([r.NombreVendedor.split(" (")[0], r.ClienteNombre[:26], money(r.PromedioMensual), money(r.EsperadoAlDia), money(r.CompradoSep), pct(r.PctCumpl)])
        cc[(i,5)] = RCLARO if r.PctCumpl<50 else (ACLARO if r.PctCumpl<100 else GCLARO)
    tabla(sl, 0.5, 1.5, 12.35, rows, [2.2,3.3,1.7,1.7,1.6,1.0], aligns=['l','l','r','r','r','c'], rowh=0.36, headh=0.5, fs=9, hfs=8.3, cell_colors=cc)
    tsum = b["topN"]
    lectura(sl, f"Top {len(tsum)}: RD$ {money(tsum.PromedioMensual.sum())} de promedio mensual, RD$ {money(tsum.EsperadoAlDia.sum())} esperado a la fecha vs. RD$ {money(tsum.CompradoSep.sum())} comprado ({pct(tsum.CompradoSep.sum()/tsum.EsperadoAlDia.sum()*100)}).", top=6.35)

    # SLIDE 4 - No visitados
    sl = prs.slides.add_slide(BLANK)
    head(sl, "3. Clientes no visitados en septiembre", "Cartera activa vs. clientes con al menos una visita registrada del 1 al 19 de septiembre.", 4, PIE)
    kpi(sl, 0.5, 1.35, 2.95, str(b["tot_cartera"]), "Cartera activa total", AZUL)
    kpi(sl, 3.62, 1.35, 2.95, str(b["tot_visitados"]), "Clientes visitados", VERDE)
    pctnv = b["tot_novis"]/b["tot_cartera"]*100
    kpi(sl, 6.74, 1.35, 2.95, f"{b['tot_novis']}  ({pct(pctnv)})", "Clientes SIN visitar", ROJO if pctnv>20 else (AMBAR if pctnv>10 else VERDE))
    kpi(sl, 9.86, 1.35, 2.99, str(b["of_total"]), "Clientes activados oferta 3 Bolones+1 Tipitin", VERDE)
    sl.shapes.add_picture(str(IMG/f"{key}_novis_vend.png"), Inches(0.5), Inches(2.55), width=Inches(12.35))
    top_nv = b["no_vis"].head(5)
    if len(top_nv):
        txt = "Mayor valor sin visitar: " + "; ".join(f"{r.CliNombre} (RD$ {money(r.PromedioMensual)}/mes, {r.RepNombre.split(' (')[0]})" for r in top_nv.itertuples())
        lectura(sl, txt[:220], color=RCLARO, barra=ROJO)

    # SLIDE 5 - Plan de accion y cierre
    sl = prs.slides.add_slide(BLANK)
    head(sl, "4-5. Plan de acci\u00f3n y compromisos de la semana", "Prioridades para lo que resta de septiembre.", 5, PIE)
    peor_vend = b["por_vend"].sort_values("PctAlcance").iloc[0]
    n_visita = min(10, len(b["no_vis"]))
    monto_visita = b["no_vis"].head(n_visita).PromedioMensual.sum()
    puntos = [
        ("1","Cerrar la brecha de presupuesto", f"{peor_vend.RepNombre} es el vendedor con menor alcance ({pct(peor_vend.PctAlcance)}). Acompa\u00f1amiento del supervisor esta semana.", ROJO if peor_vend.PctAlcance<80 else AMBAR),
        ("2","Visita especial a clientes de valor sin atender", f"{n_visita} clientes no visitados suman RD$ {money(monto_visita)} de promedio mensual. Meta: visitarlos todos antes del cierre.", ROJO),
        ("3","Activar la oferta 3 Bolones + 1 Tipitin", f"{b['of_total']} clientes activados a la fecha. Sumar activaciones nuevas usando la lista de clientes de valor.", VERDE),
        ("4","Sostener el ritmo en el resto de la ruta", f"Alcance general del equipo: {pct(t.PctAlcance)}. Meta de cierre: proyecci\u00f3n de RD$ {money(t.Proyectado)} ({pct(t.PctProyVsMes)} del mes).", AZUL),
    ]
    y=1.5
    for num,tit,txt,c in puntos:
        rect(sl, 0.5, y, 12.35, 1.05, fill=FONDO); rect(sl, 0.5, y, 0.055, 1.05, fill=c)
        rect(sl, 0.68, y+0.14, 0.38, 0.38, fill=c, radius=0.5)
        tb(sl, 0.68, y+0.22, 0.38, 0.26, num, size=11, bold=True, color=BLANCO, align=PP_ALIGN.CENTER)
        tb(sl, 1.26, y+0.10, 11.4, 0.3, tit, size=12, bold=True, color=c)
        tb(sl, 1.26, y+0.44, 11.4, 0.55, txt, size=9.5, color=GRIS, space=0)
        y += 1.16
    lectura(sl, "Compromiso: revisar estos 4 puntos en la pr\u00f3xima reuni\u00f3n con evidencia de visita y venta generada.", top=6.05)

    dest = DESTDIR / f"Reunion_{key.capitalize()}_Sep2026.pptx"
    prs.save(str(dest))
    print("OK ->", dest)

if __name__ == "__main__":
    keys = sys.argv[1:] or list(SUPERVISORES.keys())
    for k in keys:
        build(k)
