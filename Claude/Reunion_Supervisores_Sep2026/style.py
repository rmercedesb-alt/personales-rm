# -*- coding: utf-8 -*-
"""Estilo AM Comercial compartido para PPTX (colores, helpers de layout)."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

SC = Path(__file__).parent
LOGO = SC / "logo_a.png"

AZUL = RGBColor(0x20, 0x3F, 0x99); VERDE = RGBColor(0x0C, 0x94, 0x44)
GRIS = RGBColor(0x3A, 0x3A, 0x3A); GRIS2 = RGBColor(0x6E, 0x6E, 0x6E)
FONDO = RGBColor(0xF2, 0xF5, 0xF3); VCLARO = RGBColor(0xE8, 0xF5, 0xEC)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF); ROJO = RGBColor(0xC0, 0x39, 0x2B)
LINEA = RGBColor(0xD5, 0xDA, 0xD7); AMBAR = RGBColor(0xE6, 0x7E, 0x22)
RCLARO = RGBColor(0xFA, 0xE5, 0xE3); ACLARO = RGBColor(0xFD, 0xF2, 0xE4)
GCLARO = RGBColor(0xE3, 0xF3, 0xE9)


def nueva_presentacion():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    return prs


def tb(slide, l, t, w, h, text, size=11, bold=False, color=GRIS,
       align=PP_ALIGN.LEFT, space=2, anchor=MSO_ANCHOR.TOP, italic=False):
    s = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    f = s.text_frame
    f.word_wrap = True
    f.margin_left = 0; f.margin_right = 0; f.margin_top = 0; f.margin_bottom = 0
    f.vertical_anchor = anchor
    for i, ln in enumerate(str(text).split("\n")):
        p = f.paragraphs[0] if i == 0 else f.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
        r.font.color.rgb = color; r.font.name = "Calibri"
    return s


def rect(slide, l, t, w, h, fill=None, line=None, radius=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if radius is not None:
        try: shp.adjustments[0] = radius
        except Exception: pass
    if fill is None: shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None: shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if shp.has_text_frame: shp.text_frame.text = ""
    return shp


def head(slide, titulo, bajada, num, pie="AM Comercial · Confidencial · Fuente: MovilBusiness · Reunión semanal de supervisores"):
    tb(slide, 0.5, 0.30, 10.4, 0.55, titulo, size=20, bold=True, color=AZUL)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(11.62), Inches(0.32), height=Inches(0.46))
    rect(slide, 0.5, 0.92, 12.35, 0.045, fill=AZUL)
    rect(slide, 0.5, 0.92, 1.6, 0.045, fill=VERDE)
    tb(slide, 0.5, 1.03, 12.4, 0.36, bajada, size=10, color=GRIS2)
    rect(slide, 0.5, 7.12, 12.35, 0.012, fill=LINEA)
    tb(slide, 0.5, 7.20, 10.5, 0.28, pie, size=8, color=GRIS2)
    tb(slide, 11.8, 7.20, 1.05, 0.28, str(num), size=8, color=GRIS2, align=PP_ALIGN.RIGHT)


def lectura(slide, txt, top=6.55, color=VCLARO, barra=VERDE, size=10.5):
    rect(slide, 0.5, top, 12.35, 0.48, fill=color)
    rect(slide, 0.5, top, 0.06, 0.48, fill=barra)
    tb(slide, 0.72, top + 0.06, 12.0, 0.40, txt, size=size, bold=True, color=GRIS)


def kpi(slide, l, t, w, valor, etiq, color=AZUL, h=0.92):
    rect(slide, l, t, w, h, fill=FONDO)
    rect(slide, l, t, 0.055, h, fill=color)
    tb(slide, l + 0.2, t + 0.10, w - 0.3, 0.42, valor, size=18, bold=True, color=color)
    tb(slide, l + 0.2, t + 0.55, w - 0.3, 0.3, etiq, size=8.5, color=GRIS2)


def tabla(slide, l, t, w, rows, colw, head_fill=AZUL, rowh=0.285, headh=0.44,
          fs=9, hfs=8.5, aligns=None, bold_rows=(), cell_colors=None, fill_rows=None):
    tot = sum(colw); colw = [c * w / tot for c in colw]
    y = t
    for i, row in enumerate(rows):
        h = headh if i == 0 else rowh
        x = l
        if i == 0:
            rect(slide, l, y, w, h, fill=head_fill)
        elif fill_rows and i in fill_rows:
            rect(slide, l, y, w, h, fill=fill_rows[i])
        elif i % 2 == 0:
            rect(slide, l, y, w, h, fill=RGBColor(0xF7, 0xF9, 0xF8))
        for j, cell in enumerate(row):
            if cell_colors and (i, j) in cell_colors:
                rect(slide, x + 0.02, y + 0.02, colw[j] - 0.04, h - 0.04, fill=cell_colors[(i, j)])
            a = 'l' if aligns is None else aligns[j]
            al = PP_ALIGN.LEFT if a == 'l' else (PP_ALIGN.RIGHT if a == 'r' else PP_ALIGN.CENTER)
            nlines = str(cell).count("\n") + 1
            th = 0.20 * nlines
            tb(slide, x + 0.10, y + (h - th) / 2, colw[j] - 0.20, th, str(cell),
               size=hfs if i == 0 else fs, bold=(i == 0 or i in bold_rows),
               color=BLANCO if i == 0 else GRIS, align=al, space=0)
            x += colw[j]
        if i > 0:
            rect(slide, l, y + h, w, 0.008, fill=RGBColor(0xE6, 0xEA, 0xE8))
        y += h
    return y


def money(x):
    return "{:,.0f}".format(x)

def pct(x):
    return "{:.0f}%".format(x) if x == x else "n/d"
