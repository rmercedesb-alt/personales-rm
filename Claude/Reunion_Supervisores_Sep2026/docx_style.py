# -*- coding: utf-8 -*-
"""Helpers de estilo AM Comercial para informes Word (python-docx)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x20, 0x3F, 0x99)
VERDE = RGBColor(0x0C, 0x94, 0x44)
GRIS = RGBColor(0x3A, 0x3A, 0x3A)
GRIS2 = RGBColor(0x6E, 0x6E, 0x6E)
ROJO = RGBColor(0xC0, 0x39, 0x2B)
AMBAR = RGBColor(0xE6, 0x7E, 0x22)


def nuevo_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = GRIS
    for section in doc.sections:
        section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
        section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.5)
    return doc


def portada(doc, titulo, subtitulo, meta):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AM COMERCIAL"); r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = AZUL
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(titulo); r2.font.size = Pt(22); r2.font.bold = True; r2.font.color.rgb = AZUL
    p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(subtitulo); r3.font.size = Pt(13); r3.font.color.rgb = GRIS2
    p4 = doc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p4.add_run(meta); r4.font.size = Pt(10.5); r4.font.italic = True; r4.font.color.rgb = GRIS2
    doc.add_paragraph()


def h1(doc, text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = AZUL; r.font.size = Pt(16); r.font.bold = True
    return p


def h2(doc, text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = VERDE; r.font.size = Pt(12.5); r.font.bold = True
    return p


def parrafo(doc, text, size=10.5, bold=False, color=GRIS, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.italic = italic
    return p


def callout(doc, text, color=AZUL):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run("▶ " + text)
    r.font.size = Pt(10.5); r.font.bold = True; r.font.color.rgb = color
    return p


def _shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def tabla(doc, headers, rows, col_widths=None, cell_shade=None, bold_last_row=False):
    """rows: list of lists (str). cell_shade: dict {(row_idx_incl_header, col_idx): hex_color}"""
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for j, htext in enumerate(headers):
        hdr[j].text = ""
        p = hdr[j].paragraphs[0]; r = p.add_run(htext)
        r.font.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade_cell(hdr[j], "203F99")
    for i, row in enumerate(rows, start=1):
        cells = t.add_row().cells
        for j, val in enumerate(row):
            cells[j].text = ""
            p = cells[j].paragraphs[0]; r = p.add_run(str(val))
            r.font.size = Pt(9.5); r.font.color.rgb = GRIS
            bold_row = bold_last_row and i == len(rows)
            r.font.bold = bold_row
            if bold_row: _shade_cell(cells[j], "E9EDF7")
            if cell_shade and (i, j) in cell_shade:
                _shade_cell(cells[j], cell_shade[(i, j)])
    if col_widths:
        for row in t.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Cm(w)
    return t
