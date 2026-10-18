# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
import os,data as D
_B=os.path.dirname(os.path.abspath(__file__))
VERDE="0C9444"; AZUL="203F99"; GRIS="6E6E6E"; CARBON="3A3A3A"; AZCLARO="E7EBF5"; WHITE="FFFFFF"
RD='"RD$"#,##0;("RD$"#,##0);-'; PCT='0.0%'; NUM='#,##0'
OUT=os.path.dirname(_B)+"/"; CH=_B+"/charts/"; DIAS=D.DIAS_HABILES_JUL
thin=Side(style="thin",color="D9D9D9"); border=Border(left=thin,right=thin,top=thin,bottom=thin)
def A(sz=10,b=False,color="000000"): return Font(name="Arial",size=sz,bold=b,color=color)
def fill(c): return PatternFill("solid",fgColor=c)
def hdr(cell,txt,bg=VERDE,fg=WHITE):
    cell.value=txt; cell.font=A(b=True,color=fg); cell.fill=fill(bg)
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=border
def cval(cell,v,fmt=None,b=False,color="000000",align="right",bg=None):
    cell.value=v; cell.font=A(b=b,color=color)
    if fmt:cell.number_format=fmt
    cell.alignment=Alignment(horizontal=align,vertical="center"); cell.border=border
    if bg:cell.fill=fill(bg)
def sem(p): return "🟢" if p>=1.0 else ("🟡" if p>=0.85 else "🔴")
wb=openpyxl.load_workbook(OUT+f"AM_Cierre_{D.MES_NOMBRE}2026.xlsx")

def bloque(centro,titulo,color,grupos,provkey,abcdkey,chart_prov,chart_gauss,chart_abcd):
    ws=wb.create_sheet(titulo); ws.sheet_view.showGridLines=False
    for col,w in zip("ABCDEFGHI",[3,32,13,13,12,12,12,11,10]): ws.column_dimensions[col].width=w
    ws["B1"].value=f"{titulo} — Cierre {D.MES_NOMBRE} 2026 (neto, ex-Caobo)"; ws["B1"].font=A(sz=12,b=True,color=CARBON)
    intro_txt=D.NARR["CAPS"][centro]["intro"]
    ws["B2"].value=intro_txt; ws["B2"].font=A(sz=9,color=GRIS)
    venta=D.serie_2026[D.MES][0] if centro=="Santo Domingo" else D.serie_2026[D.MES][1]
    io=D.ind_op[centro]; facts=sum(p[3] for p in D.prov_jul[provkey]); ticket=venta/facts
    r=3
    ws.cell(r,2).value="1 · Indicadores de operación"; ws.cell(r,2).font=A(sz=11,b=True,color=color); r+=1
    ind=[(f"Venta neta {D.MES_NOMBRE}","RD$",venta,RD),(f"Productividad por día hábil (÷{D.DIAS_HABILES})","RD$",venta/DIAS,RD),
         ("Ticket promedio (venta ÷ facturas)","RD$",ticket,RD),
         ("Clientes activos (facturados)","N°",io["activos"],NUM),
         ("Cartera (clientes en ruta, distinta)","N°",io["cartera"],NUM),
         ("Activación de cartera (activos ÷ cartera)","%",io["activacion"],PCT),
         ("Efectividad de visita (con pedido ÷ realizadas)","%",io["efectividad"],PCT)]
    for j,h in enumerate(["Indicador","Unidad",f"Valor {D.MES_ABBR} 2026"]): hdr(ws.cell(r,2+j),h)
    for i,(nm,u,v,fmt) in enumerate(ind):
        rr=r+1+i; cval(ws.cell(rr,2),nm,align="left"); cval(ws.cell(rr,3),u,align="center"); cval(ws.cell(rr,4),v,fmt)
    try:
        im=XLImage(CH+"operacion.png"); im.width=int(im.width*0.52); im.height=int(im.height*0.52); ws.add_image(im,"F3")
    except Exception as e:print(e)
    r=r+len(ind)+2
    ws.cell(r,2).value="2 · Comportamiento de proveedores (Top 4 + Otros)"; ws.cell(r,2).font=A(sz=11,b=True,color=color); r+=1
    for j,h in enumerate(["Proveedor","Neto Jul","% part.","Clientes","Facturas","Ticket","Prod/cliente"]): hdr(ws.cell(r,2+j),h)
    lst=D.prov_jul[provkey]; top=lst[:4]; otros=lst[4:]; tot_neto=sum(p[1] for p in lst)
    r0=r+1; rows2=list(top)+[("Otros",sum(p[1] for p in otros),sum(p[2] for p in otros),sum(p[3] for p in otros))]
    for i,(nm,ne,cl,fa) in enumerate(rows2):
        rr=r0+i; isot=(nm=="Otros")
        cval(ws.cell(rr,2),nm,align="left",b=(not isot),color=(GRIS if isot else "000000"))
        cval(ws.cell(rr,3),ne,RD); cval(ws.cell(rr,4),ne/tot_neto,PCT); cval(ws.cell(rr,5),cl,NUM); cval(ws.cell(rr,6),fa,NUM)
        cval(ws.cell(rr,7),ne/fa,RD); cval(ws.cell(rr,8),ne/cl,RD)
    rr=r0+len(rows2); cval(ws.cell(rr,2),"TOTAL centro",align="left",b=True,bg=AZCLARO)
    cval(ws.cell(rr,3),f"=SUM(C{r0}:C{rr-1})",RD,b=True,bg=AZCLARO); cval(ws.cell(rr,4),f"=SUM(D{r0}:D{rr-1})",PCT,b=True,bg=AZCLARO)
    for cc in (5,6): cval(ws.cell(rr,cc),f"=SUM({get_column_letter(cc)}{r0}:{get_column_letter(cc)}{rr-1})",NUM,b=True,bg=AZCLARO)
    cval(ws.cell(rr,7),"",bg=AZCLARO); cval(ws.cell(rr,8),"",bg=AZCLARO)
    try:
        im=XLImage(CH+chart_prov); im.width=int(im.width*0.58); im.height=int(im.height*0.58); ws.add_image(im,f"J{r}")
    except Exception as e:print(e)
    r=rr+2
    ws.cell(r,2).value=f"3 · Clientes A·B·C·D ({D.MES_NOMBRE} 2026)"; ws.cell(r,2).font=A(sz=11,b=True,color=color); r+=1
    for j,h in enumerate(["Segmento","Neto Jul","% part.","Clientes","Prod/cliente"]): hdr(ws.cell(r,2+j),h)
    segs=[("A Grande","A"),("B Mediano","B"),("C Pequeño","C"),("D Muy pequeño","D"),("Sin clasif.","SC")]
    dd=D.abcd[abcdkey][D.MES]; totn=sum(dd[s][0] for s in [x[1] for x in segs]); r0=r+1
    for i,(lbl,s) in enumerate(segs):
        rr=r0+i; ne,cl=dd[s]
        cval(ws.cell(rr,2),lbl,align="left"); cval(ws.cell(rr,3),ne,RD); cval(ws.cell(rr,4),ne/totn,PCT); cval(ws.cell(rr,5),cl,NUM); cval(ws.cell(rr,6),ne/cl if cl else 0,RD)
    rr=r0+len(segs); cval(ws.cell(rr,2),"TOTAL",align="left",b=True,bg=AZCLARO)
    cval(ws.cell(rr,3),f"=SUM(C{r0}:C{rr-1})",RD,b=True,bg=AZCLARO); cval(ws.cell(rr,4),f"=SUM(D{r0}:D{rr-1})",PCT,b=True,bg=AZCLARO)
    cval(ws.cell(rr,5),f"=SUM(E{r0}:E{rr-1})",NUM,b=True,bg=AZCLARO); cval(ws.cell(rr,6),"",bg=AZCLARO)
    try:
        im=XLImage(CH+chart_abcd); im.width=int(im.width*0.58); im.height=int(im.height*0.58); ws.add_image(im,f"H{r}")
    except Exception as e:print(e)
    r=rr+2
    ws.cell(r,2).value="4 · Fuerza de venta — Cumplimiento de cuota + Efectividad (campana)"; ws.cell(r,2).font=A(sz=11,b=True,color=color); r+=1
    grpname=D.GRPNAME
    for j,h in enumerate(["Grupo/FFVV","Vendedor","Cuota Jul","Ejecutado","% Cuota","Sem.","Efectiv."]): hdr(ws.cell(r,2+j),h)
    vs=[v for v in D.vendedores if v[1]==centro and v[2] in grupos]
    vs=sorted(vs,key=lambda v:-(v[5]/v[4] if v[4]>0 else 99)); r0=r+1
    for i,v in enumerate(vs):
        rr=r0+i; cod,ce,gr,nm,cu,ej=v; pc=(ej/cu if cu>0 else None)
        cval(ws.cell(rr,2),grpname.get(gr,gr),align="left"); cval(ws.cell(rr,3),nm,align="left")
        cval(ws.cell(rr,4),cu,RD); cval(ws.cell(rr,5),ej,RD)
        if pc is None: cval(ws.cell(rr,6),"s/cuota",color=GRIS); cval(ws.cell(rr,7),"—",align="center")
        else: cval(ws.cell(rr,6),f"=E{rr}/D{rr}",PCT); cval(ws.cell(rr,7),sem(pc),align="center")
        ef=D.efect.get(cod); cval(ws.cell(rr,8),(ef[1]/ef[0] if ef and ef[0] else None),PCT) if ef else cval(ws.cell(rr,8),"—",align="center")
    rr=r0+len(vs); cval(ws.cell(rr,2),"TOTAL centro",align="left",b=True,bg=AZCLARO); cval(ws.cell(rr,3),"",bg=AZCLARO)
    cval(ws.cell(rr,4),f"=SUM(D{r0}:D{rr-1})",RD,b=True,bg=AZCLARO); cval(ws.cell(rr,5),f"=SUM(E{r0}:E{rr-1})",RD,b=True,bg=AZCLARO)
    cval(ws.cell(rr,6),f"=E{rr}/D{rr}",PCT,b=True,bg=AZCLARO); cval(ws.cell(rr,7),f'=IF(F{rr}>=1,"🟢",IF(F{rr}>=0.85,"🟡","🔴"))',align="center",bg=AZCLARO)
    cval(ws.cell(rr,8),io["efectividad"],PCT,b=True,bg=AZCLARO)
    try:
        im=XLImage(CH+chart_gauss); im.width=int(im.width*0.6); im.height=int(im.height*0.6); ws.add_image(im,f"J{r}")
    except Exception as e:print(e)

bloque("Santo Domingo","Bloque A - SDO",AZUL,D.GRUPOS_SDO,"Santo Domingo","Santo Domingo","prov_sdo.png","gauss_sdo.png","abcd_sdo.png")
bloque("Azua","Bloque B - Sur (Azua)",VERDE,D.GRUPOS_AZUA,"Azua","Azua","prov_azua.png","gauss_azua.png","abcd_azua.png")

# FFVV split resumen (Azua)
wsz=wb["Bloque B - Sur (Azua)"]; rr=3
wsz.cell(rr,16).value="FFVV Azua"; wsz.cell(rr,16).font=A(sz=10,b=True,color=VERDE)
for j,h in enumerate(["Fuerza","Vend.","Ejec. Jul","%centro"]): hdr(wsz.cell(rr+1,16+j),h)
grp_tot={}
for v in D.vendedores:
    if v[1]=="Azua": grp_tot.setdefault(v[2],[0,0]); grp_tot[v[2]][0]+=1; grp_tot[v[2]][1]+=v[5]
tot_az=sum(g[1] for g in grp_tot.values())
labels=dict(D.FFVV_AZUA)
for i,g in enumerate([gg for gg,_ in D.FFVV_AZUA]):
    r2=rr+2+i; n,ej=grp_tot[g]
    cval(wsz.cell(r2,16),labels[g],align="left"); cval(wsz.cell(r2,17),n,NUM); cval(wsz.cell(r2,18),ej,RD); cval(wsz.cell(r2,19),ej/tot_az,PCT)
for col,w in zip(["P","Q","R","S"],[22,7,14,9]): wsz.column_dimensions[col].width=w

# COLOMBINA
ws=wb.create_sheet("Apéndice Colombina"); ws.sheet_view.showGridLines=False
for col,w in zip("ABCDEF",[3,32,14,14,10,12]): ws.column_dimensions[col].width=w
ws["B1"].value=f"Apéndice · Colombina — Top 5 SKU por CEDIS ({D.MES_NOMBRE} 2026)"; ws["B1"].font=A(sz=12,b=True,color=CARBON)
ws["B2"].value="Colombina se apoya en BBB/Grissly/Tipitin; galletería salió del portafolio (mar-2026). Riesgo de concentración en pocos SKU."
ws["B2"].font=A(sz=9,color=GRIS)
r=4
for centro,dataset,tot in [("Santo Domingo",D.colombina_sdo,D.colombina_total_sdo),("Azua",D.colombina_azua,D.colombina_total_azua)]:
    ws.cell(r,2).value=f"{centro} — Colombina total RD${tot:,.0f}"; ws.cell(r,2).font=A(sz=11,b=True,color=(AZUL if centro=="Santo Domingo" else VERDE)); r+=1
    for j,h in enumerate(["SKU","Neto Jul","% de Colombina","Clientes","Prod/cliente"]): hdr(ws.cell(r,2+j),h)
    r0=r+1
    for i,(sku,ne,cl) in enumerate(dataset):
        rr=r0+i; cval(ws.cell(rr,2),sku,align="left"); cval(ws.cell(rr,3),ne,RD); cval(ws.cell(rr,4),ne/tot,PCT); cval(ws.cell(rr,5),cl,NUM); cval(ws.cell(rr,6),ne/cl,RD)
    rr=r0+len(dataset)
    cval(ws.cell(rr,2),"Top 5 subtotal",align="left",b=True,bg=AZCLARO); cval(ws.cell(rr,3),f"=SUM(C{r0}:C{rr-1})",RD,b=True,bg=AZCLARO)
    cval(ws.cell(rr,4),f"=SUM(D{r0}:D{rr-1})",PCT,b=True,bg=AZCLARO); cval(ws.cell(rr,5),f"=SUM(E{r0}:E{rr-1})",NUM,b=True,bg=AZCLARO); cval(ws.cell(rr,6),"",bg=AZCLARO)
    r=rr+2
try:
    im=XLImage(CH+"colombina.png"); im.width=int(im.width*0.72); im.height=int(im.height*0.72); ws.add_image(im,"H4")
except Exception as e:print(e)

# CIERRE
ws=wb.create_sheet("Cierre - Acciones"); ws.sheet_view.showGridLines=False
for col,w in zip("ABCDEF",[3,18,42,36,20,10]): ws.column_dimensions[col].width=w
ws["B1"].value="Cierre · Conclusiones y recomendaciones accionables"; ws["B1"].font=A(sz=12,b=True,color=CARBON)
r=3
for j,h in enumerate(["Señal","Qué pasó / Dónde","Qué hacer","Responsable","Fecha"]): hdr(ws.cell(r,2+j),h)
def _emoji(t): return t.replace("{G}","🟢").replace("{Y}","🟡").replace("{R}","🔴")
acc=[(_emoji(a),_emoji(b),c,d,e) for a,b,c,d,e in D.NARR["CIERRE"]]
r0=r+1
for i,(s,q,h,resp,fe) in enumerate(acc):
    rr=r0+i
    cval(ws.cell(rr,2),s,align="left",b=True); cval(ws.cell(rr,3),q,align="left"); cval(ws.cell(rr,4),h,align="left"); cval(ws.cell(rr,5),resp,align="left"); cval(ws.cell(rr,6),fe,align="center")
    for cc in range(2,7): ws.cell(rr,cc).alignment=Alignment(wrap_text=True,vertical="top",horizontal=ws.cell(rr,cc).alignment.horizontal)
    ws.row_dimensions[rr].height=52
ws.cell(r0+len(acc)+1,2).value="AM Comercial — Confidencial"; ws.cell(r0+len(acc)+1,2).font=A(sz=8,color=GRIS)

# Hoja Caobo por suplidor (apéndice para Presidencia)
if getattr(D,"caobo_prov",None):
    ws=wb.create_sheet("Ventas a Caobo"); ws.sheet_view.showGridLines=False
    for col,w in zip("ABCDEF",[3,26,15,15,15,12]): ws.column_dimensions[col].width=w
    ws["B1"].value=f"Apéndice · Ventas a Inversiones Caobo por suplidor ({D.MES_NOMBRE} 2026)"; ws["B1"].font=A(sz=12,b=True,color=CARBON)
    ws["B2"].value=D.NARR.get("CAOBO_INTRO",""); ws["B2"].font=A(sz=9,color=GRIS)
    r=4
    for j,h in enumerate(["Suplidor","Neto 2026","Con ITBIS 2026","Neto 2025","Δ% interan."]): hdr(ws.cell(r,2+j),h)
    cp=sorted(D.caobo_prov,key=lambda x:-x[1]); r0=r+1
    for i,(nm,n,it,n25) in enumerate(cp):
        rr=r0+i; cval(ws.cell(rr,2),nm,align="left"); cval(ws.cell(rr,3),n,RD); cval(ws.cell(rr,4),it,RD); cval(ws.cell(rr,5),n25,RD)
        cval(ws.cell(rr,6),(n/n25-1) if n25 else 0,PCT,color=(VERDE if (n25 and n>=n25) else "C0392B"))
    rr=r0+len(cp)
    cval(ws.cell(rr,2),"TOTAL",align="left",b=True,bg=AZCLARO)
    cval(ws.cell(rr,3),f"=SUM(C{r0}:C{rr-1})",RD,b=True,bg=AZCLARO); cval(ws.cell(rr,4),f"=SUM(D{r0}:D{rr-1})",RD,b=True,bg=AZCLARO)
    cval(ws.cell(rr,5),f"=SUM(E{r0}:E{rr-1})",RD,b=True,bg=AZCLARO); cval(ws.cell(rr,6),f"=C{rr}/E{rr}-1",PCT,b=True,bg=AZCLARO)
    try:
        im=XLImage(CH+"caobo.png"); im.width=int(im.width*0.6); im.height=int(im.height*0.6); ws.add_image(im,"H4")
    except Exception as e: print(e)
order=["Resumen","Capa 1 - General","Bloque A - SDO","Bloque B - Sur (Azua)","Apéndice Colombina","Cierre - Acciones","Ventas a Caobo"]
wb._sheets.sort(key=lambda s: order.index(s.title) if s.title in order else 99)
wb.save(OUT+f"AM_Cierre_{D.MES_NOMBRE}2026.xlsx")
print("Excel completo v2 guardado")
