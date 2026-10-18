"""
Sube los datos del archivo CARGA_REAL_PRUEBAS.xlsx a Google Sheets.
Ejecutar: C:/Python314/python.exe subir_a_sheets.py
"""
import openpyxl, warnings, time, sys, io
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

warnings.filterwarnings('ignore')
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

# ── Configuración ────────────────────────────────────────────────────────────
SPREADSHEET_ID  = '1Ir6aKzaL8HRSi0v8lUMa_1KYTQc0r4gtTVwHiHiYbAU'
CREDENTIALS_PATH = r'C:\Personales RM\Claude\Sistema de Condominios\web-app\InmueblesAPI\credentials.json'
EXCEL_PATH       = r'C:\Personales RM\Claude\Sistema de Condominios\web-app\CARGA_REAL_PRUEBAS.xlsx'

# Mapeo: pestaña Excel → pestaña Google Sheets
MAPEO = {
    '🏠 PROPIETARIOS':   '👤 PROPIETARIOS',
    '🏢 INMUEBLES':      '🏢 INMUEBLES',
    '👥 INQUILINOS':     '👥 INQUILINOS',
    '📋 CONTRATOS':      '📋 CONTRATOS',
    '💰 CXC_PENDIENTES': '💰 CXC_CARGOS',
    '✅ PAGOS_HISTORICOS':'💳 CXC_PAGOS',
}

BATCH_SIZE = 50   # filas por petición (evita timeout)
PAUSA      = 1.0  # segundos entre batches (respetar rate limit de Sheets API)

# ── Autenticación ────────────────────────────────────────────────────────────
print('Conectando con Google Sheets...')
creds = service_account.Credentials.from_service_account_file(
    CREDENTIALS_PATH,
    scopes=['https://www.googleapis.com/auth/spreadsheets']
)
service = build('sheets', 'v4', credentials=creds)
sheets  = service.spreadsheets()
print('✓ Conectado\n')

# ── Leer Excel ───────────────────────────────────────────────────────────────
print(f'Leyendo {EXCEL_PATH}...')
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
print('✓ Archivo leído\n')

def limpiar_valor(v):
    """Convierte valores Python a string limpio para Sheets."""
    if v is None: return ''
    import datetime
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    return '' if s in ('None', 'nan') else s

def subir_hoja(tab_excel: str, tab_sheets: str):
    if tab_excel not in wb.sheetnames:
        print(f'  ⚠ Pestaña "{tab_excel}" no encontrada en el Excel, saltando.')
        return

    ws = wb[tab_excel]
    filas = []
    for row in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        # Saltar filas completamente vacías
        if not any(v for v in row if v is not None and str(v).strip()):
            continue
        # Saltar filas donde la columna A (ID) está vacía
        if not row[0] or not str(row[0]).strip():
            continue
        filas.append([limpiar_valor(v) for v in row])

    if not filas:
        print(f'  ⚠ No hay datos en "{tab_excel}", saltando.')
        return

    print(f'  Subiendo {len(filas)} filas a "{tab_sheets}"...')

    # Subir en batches
    subidas = 0
    for i in range(0, len(filas), BATCH_SIZE):
        batch = filas[i:i + BATCH_SIZE]
        body  = {'values': batch}
        try:
            sheets.values().append(
                spreadsheetId=SPREADSHEET_ID,
                range=f"'{tab_sheets}'!A:A",
                valueInputOption='USER_ENTERED',
                insertDataOption='INSERT_ROWS',
                body=body
            ).execute()
            subidas += len(batch)
            pct = int(subidas / len(filas) * 100)
            print(f'    {pct}% ({subidas}/{len(filas)})', end='\r')
            time.sleep(PAUSA)
        except HttpError as e:
            print(f'\n  ✗ Error en batch {i}-{i+len(batch)}: {e}')
            raise

    print(f'    ✓ {len(filas)} filas subidas a "{tab_sheets}"            ')

# ── Verificar que las pestañas existen en Sheets ─────────────────────────────
print('Verificando pestañas en Google Sheets...')
meta = sheets.get(spreadsheetId=SPREADSHEET_ID).execute()
tabs_existentes = {s['properties']['title'] for s in meta['sheets']}

pestañas_faltantes = [t for t in MAPEO.values() if t not in tabs_existentes]
if pestañas_faltantes:
    print('\n⚠ Las siguientes pestañas NO existen en tu Google Sheets:')
    for t in pestañas_faltantes:
        print(f'   → "{t}"')
    print('\nDebes crearlas manualmente en Google Sheets con esos nombres exactos')
    print('(incluyendo el emoji) antes de continuar.')
    print('\nPresiona Enter si ya las creaste, o Ctrl+C para cancelar...')
    input()
else:
    print('✓ Todas las pestañas existen\n')

# ── Ejecutar carga ───────────────────────────────────────────────────────────
print('=' * 55)
print('INICIANDO CARGA DE DATOS')
print('=' * 55)

errores = []
for tab_excel, tab_sheets in MAPEO.items():
    print(f'\n[{tab_excel}] → [{tab_sheets}]')
    try:
        subir_hoja(tab_excel, tab_sheets)
    except Exception as e:
        errores.append((tab_excel, str(e)))
        print(f'  ✗ FALLÓ: {e}')

print('\n' + '=' * 55)
if errores:
    print(f'CARGA COMPLETADA CON {len(errores)} ERROR(ES):')
    for tab, err in errores:
        print(f'  ✗ {tab}: {err}')
else:
    print('✓ CARGA COMPLETADA SIN ERRORES')
    print(f'\nResumen:')
    totales = {
        'Propietarios':  len(wb['🏠 PROPIETARIOS'].values) - 2 if '🏠 PROPIETARIOS' in wb.sheetnames else 0,
        'Inmuebles':     len(wb['🏢 INMUEBLES'].values) - 2 if '🏢 INMUEBLES' in wb.sheetnames else 0,
        'Inquilinos':    len(wb['👥 INQUILINOS'].values) - 2 if '👥 INQUILINOS' in wb.sheetnames else 0,
        'Contratos':     len(wb['📋 CONTRATOS'].values) - 2 if '📋 CONTRATOS' in wb.sheetnames else 0,
        'CxC Pendientes':len(wb['💰 CXC_PENDIENTES'].values) - 2 if '💰 CXC_PENDIENTES' in wb.sheetnames else 0,
        'Pagos':         len(wb['✅ PAGOS_HISTORICOS'].values) - 2 if '✅ PAGOS_HISTORICOS' in wb.sheetnames else 0,
    }
    for k, v in totales.items():
        print(f'  {k:<20}: {v} registros')

print('\n✓ Puedes abrir el sistema y verificar los datos.')
print(f'  URL: https://docs.google.com/spreadsheets/d/{SPREADSHEET_ID}')
