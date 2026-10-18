"""
Genera CARGA_REAL_PRUEBAS.xlsx con los datos reales del Reporte de Inquilinos.
Correcciones aplicadas:
  - Inmueble = unidad (col A), Inquilino = persona (col B)
  - Coordenadas aleatorias en La Romana, La Altagracia y San Pedro de Macorís
  - Vacíos = inmuebles DISPONIBLES (sin inquilino ni contrato)
  - Ene-Mar pagos -> LIQUIDADO | Abr-May -> EN_ADMINISTRACION | Jun -> EN_ADMINISTRACION
  - Meses sin fecha (pasados) -> CxC PENDIENTE o EN_MORA
  - Junio: si no pagó -> CxC PENDIENTE (mes actual)
"""
import openpyxl, datetime, warnings, random, calendar
warnings.filterwarnings('ignore')

src = openpyxl.load_workbook(
    r'C:\Personales RM\Claude\Sistema de Condominios\web-app\Reporte de Inquilinos BackUp.xlsx',
    data_only=True)

HOY = datetime.date(2026, 6, 14)

# ── Coordenadas por provincia (bounding boxes reales) ───────────────────────
PROVINCIAS = {
    'La Romana':             {'lat': (18.38, 18.52), 'lng': (-69.05, -68.82)},
    'La Altagracia':         {'lat': (18.50, 18.82), 'lng': (-68.95, -68.45)},
    'San Pedro de Macorís':  {'lat': (18.40, 18.60), 'lng': (-69.48, -69.18)},
}
PROV_NAMES = list(PROVINCIAS.keys())

def rand_coords(prov):
    b = PROVINCIAS[prov]
    lat = round(random.uniform(*b['lat']), 6)
    lng = round(random.uniform(*b['lng']), 6)
    return lat, lng

# Asignar provincia a cada pestaña de forma fija y reproducible
def provincia_para_sheet(idx):
    return PROV_NAMES[idx % 3]

def limpiar_monto(v):
    if v is None: return 0
    s = str(v).replace('RD$','').replace('RD','').replace(',','') \
               .replace(' ','').replace('.00','').replace('$','')
    # Casos tipo "5,5000" (error tipográfico)
    try: return max(0.0, float(s))
    except: return 0

def es_vacio(v):
    if v is None: return True
    return str(v).lower().strip() in ('','vacio','vacia','vacant','none','--',
                                       '#value!','rd','vacio','vacia',' ')

def es_fecha_valida(v):
    if not isinstance(v, (datetime.datetime, datetime.date)): return False
    d = v.date() if isinstance(v, datetime.datetime) else v
    return 2025 <= d.year <= 2026

def to_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    return None

def fmt(d):
    if d is None: return ''
    return d.strftime('%Y-%m-%d') if isinstance(d, (datetime.date, datetime.datetime)) else ''

def custodia(fecha_pago: datetime.date) -> str:
    """Ene-Mar 2026 → LIQUIDADO, Abr en adelante → EN_ADMINISTRACION"""
    if fecha_pago.year < 2026: return 'LIQUIDADO'
    if fecha_pago.month <= 3:  return 'LIQUIDADO'
    return 'EN_ADMINISTRACION'

# ── Estructura de datos ──────────────────────────────────────────────────────
propietarios, inmuebles_disp, inmuebles_ocup = [], [], []
inquilinos, contratos = [], []
cxc_pendientes, cxc_pagadas = [], []

prop_idx = inm_idx = inq_idx = ctr_idx = cxc_idx = pag_idx = 1

# ── Meses del año que cubre el archivo: Ene-Jun 2026 ────────────────────────
MESES = [1, 2, 3, 4, 5, 6]  # columnas D a I en la mayoría de hojas

for sheet_num, sheet_name in enumerate(src.sheetnames):
    ws = src[sheet_name]
    prov = provincia_para_sheet(sheet_num)
    prop_id = f'TEST-PROP-{prop_idx:03d}'
    prop_idx += 1

    propietarios.append([
        prop_id, sheet_name.strip(), '', '809-000-0000', '',
        'BanReservas', '', 'AHORRO',
        'PORCENTAJE', 10, '', 'ACTIVO', prov, '2026-01-01'
    ])

    # Detectar offset de columnas según formato de la hoja
    # Sucesión Oliva: col A=owner, B=unit, C=nombre, D=monto, E+=fechas
    # Resto: col A=unit, B=nombre, C=monto, D+=fechas
    es_oliva = 'Oliva' in sheet_name

    filas_datos = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if not any(v for v in row if v is not None): continue

        if es_oliva:
            unidad   = str(row[1]).strip() if len(row) > 1 and row[1] else ''
            nombre   = str(row[2]).strip() if len(row) > 2 and row[2] else ''
            monto_v  = row[3] if len(row) > 3 else None
            fechas_v = list(row[4:]) if len(row) > 4 else []
        else:
            unidad   = str(row[0]).strip() if row[0] else ''
            nombre   = str(row[1]).strip() if len(row) > 1 and row[1] else ''
            monto_v  = row[2] if len(row) > 2 else None
            fechas_v = list(row[3:]) if len(row) > 3 else []

        monto = limpiar_monto(monto_v)
        if monto <= 0: continue

        # Filtrar encabezados y totales
        n_low = nombre.lower()
        if not nombre or n_low in ('nombre','nombres','monto','total','casas',
                                    'apartamentos','apartamento','local',
                                    'habitaciones','nombre dueño','|nonbre'): continue
        if unidad.lower() in ('nombre','monto','casas','total','habitaciones',
                               'apartamentos','casa ','apto'): continue

        # Recolectar fechas de pago del año 2026 por mes
        pagos_por_mes = {}
        for v in fechas_v:
            d = to_date(v)
            if d and es_fecha_valida(v) and d.year == 2026:
                m = d.month
                if m not in pagos_por_mes:
                    pagos_por_mes[m] = d

        filas_datos.append((unidad, nombre, monto, pagos_por_mes))

    # Deduplicar por nombre de inquilino (misma persona puede aparecer en sección jul-dic)
    vistos = {}
    for unidad, nombre, monto, pagos_por_mes in filas_datos:
        key = (unidad.lower()[:20], nombre.lower()[:20])
        if key in vistos:
            # Fusionar pagos
            for m, d in pagos_por_mes.items():
                if m not in vistos[key][3]:
                    vistos[key][3][m] = d
        else:
            vistos[key] = [unidad, nombre, monto, pagos_por_mes]

    filas_unicas = list(vistos.values())

    for unidad, nombre, monto, pagos_por_mes in filas_unicas:
        lat, lng = rand_coords(prov)
        inm_id = f'TEST-INM-{inm_idx:03d}'
        inm_idx += 1

        # Tipo de inmueble
        u = unidad.lower()
        if 'local' in u:            tipo = 'LOCAL'
        elif 'oficina' in u:        tipo = 'OFICINA'
        elif 'casa' in u:           tipo = 'CASA'
        elif 'apto' in u or 'apt' in u or '#' in u: tipo = 'APARTAMENTO'
        elif 'estudio' in u:        tipo = 'APARTAMENTO'
        elif 'nivel' in u:          tipo = 'APARTAMENTO'
        else:                        tipo = 'APARTAMENTO'

        desc = f"{unidad} – {sheet_name.strip()}"[:60] if unidad else sheet_name.strip()[:60]

        es_vacante = es_vacio(nombre)

        if es_vacante:
            # Inmueble disponible — sin inquilino ni contrato
            inmuebles_disp.append([
                inm_id, prop_id, tipo, desc,
                '', '', prov, lat, lng,
                monto, 'DISPONIBLE', 'GRIS', unidad, '2026-01-01'
            ])
            continue

        # Inmueble ocupado
        nivel_mora = 'VERDE'
        if 6 not in pagos_por_mes:
            if 5 not in pagos_por_mes: nivel_mora = 'NARANJA'
            else:                       nivel_mora = 'AMARILLO'

        inmuebles_ocup.append([
            inm_id, prop_id, tipo, desc,
            '', '', prov, lat, lng,
            monto, 'OCUPADO', nivel_mora, unidad, '2026-01-01'
        ])

        # Inquilino
        inq_id = f'TEST-INQ-{inq_idx:03d}'
        inq_idx += 1
        inquilinos.append([
            inq_id, nombre, '', '809-000-0000', '', '',
            '', '', 'ACTIVO', '', '2026-01-01'
        ])

        # Contrato
        ctr_id = f'TEST-CTR-{ctr_idx:03d}'
        ctr_idx += 1
        contratos.append([
            ctr_id, inm_id, inq_id,
            '2026-01-01', '2026-12-31',
            monto, 'MENSUAL', 1, 30,
            'PORCENTAJE', 5, '',
            monto, 'VIGENTE', '', '2026-01-01'
        ])

        # ── CxC y Pagos por cada mes ─────────────────────────────────────
        for mes in MESES:
            ini_mes = datetime.date(2026, mes, 1)
            ult_dia = calendar.monthrange(2026, mes)[1]
            fin_mes = datetime.date(2026, mes, ult_dia)
            venc    = datetime.date(2026, mes, 1)   # día de cobro = 1
            dias_desde_venc = (HOY - venc).days

            cxc_id = f'TEST-CXC-{cxc_idx:04d}'
            cxc_idx += 1

            if mes in pagos_por_mes:
                fecha_pago = pagos_por_mes[mes]
                pag_id = f'TEST-PAG-{pag_idx:04d}'
                pag_idx += 1
                cust = custodia(fecha_pago)

                cxc_pagadas.append([
                    pag_id, cxc_id, ctr_id, inq_id, nombre, inm_id,
                    'MIGRACION', fmt(fecha_pago),
                    monto, 0, monto, 'EFECTIVO', f'MIGR-{pag_id}',
                    cust, '' if cust != 'LIQUIDADO' else 'LIQ-MIGR',
                    f'Alquiler {ini_mes.strftime("%B %Y")}', '2026-01-01'
                ])

            else:
                # Solo registrar CxC pendientes de meses ya vencidos
                # Junio: mes actual, también se registra como pendiente
                if mes <= HOY.month:
                    dias_mora = max(0, dias_desde_venc - 30)
                    mora = round(monto * 0.05, 2) if dias_mora > 0 else 0
                    estado = 'EN_MORA' if dias_mora > 0 else 'PENDIENTE'
                    saldo  = monto + mora

                    cxc_pendientes.append([
                        cxc_id, ctr_id, inq_id, nombre, inm_id,
                        'ALQUILER', f'Alquiler {ini_mes.strftime("%B %Y")}',
                        fmt(ini_mes), fmt(fin_mes), fmt(venc),
                        monto, mora, 0, saldo, estado,
                        dias_mora, '', '', '2026-01-01'
                    ])

# ── Combinar inmuebles (disponibles primero, luego ocupados) ─────────────────
all_inmuebles = inmuebles_disp + inmuebles_ocup

# ── Crear Excel ──────────────────────────────────────────────────────────────
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

AZUL_DARK = '0E2F55'; AZUL_CORP = '1B4F8C'; GRIS = 'F5F5F5'
VERDE_H = '1B5E20';   PURP_H = '4A148C'

def hf(c): return PatternFill('solid', fgColor=c)
def fn(bold=False, color='000000', size=9, italic=False):
    return Font(name='Arial', bold=bold, color=color, size=size, italic=italic)
def al_c(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def al_l(): return Alignment(horizontal='left', vertical='center', wrap_text=False)
def brd():
    s = Side(style='thin', color='CCCCCC')
    return Border(left=s, right=s, top=s, bottom=s)

def banner(ws, text, ncols, bg=AZUL_DARK):
    last = chr(64 + ncols)
    ws.merge_cells(f'A1:{last}1')
    c = ws['A1']
    c.value = text; c.font = fn(bold=True, color='FFFFFF', size=11)
    c.fill = hf(bg); c.alignment = al_c()
    ws.row_dimensions[1].height = 30

def write_hdr(ws, row, headers, bg=AZUL_CORP):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = fn(bold=True, color='FFFFFF')
        cell.fill = hf(bg); cell.alignment = al_c(); cell.border = brd()

def write_row(ws, row, values, shade=False):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font = fn(size=9); cell.alignment = al_l(); cell.border = brd()
        if shade: cell.fill = hf(GRIS)

def set_widths(ws, cols, widths):
    for col, w in zip(cols, widths):
        ws.column_dimensions[col].width = w

# PROPIETARIOS
ws0 = wb.active; ws0.title = '🏠 PROPIETARIOS'
ws0.sheet_view.showGridLines = False
banner(ws0, f'PROPIETARIOS — {len(propietarios)} registros TEST', 14)
H = ['ID','NOMBRE','CÉDULA','TELÉFONO','EMAIL','BANCO','NUM CUENTA','TIPO CTA',
     'TIPO COMIS','COMIS %','COMIS FIJA','ESTADO','PROVINCIA','FECHA REG']
write_hdr(ws0, 2, H)
for i, r in enumerate(propietarios, 3): write_row(ws0, i, r, i%2==0)
set_widths(ws0, 'ABCDEFGHIJKLMN', [14,28,13,13,22,14,16,10,12,10,10,10,22,12])
ws0.freeze_panes = 'A3'

# INMUEBLES
ws1 = wb.create_sheet('🏢 INMUEBLES')
ws1.sheet_view.showGridLines = False
n_disp = len(inmuebles_disp); n_ocup = len(inmuebles_ocup)
banner(ws1, f'INMUEBLES — {len(all_inmuebles)} unidades ({n_ocup} ocupadas, {n_disp} disponibles) — coordenadas editables en el sistema', 14)
H1 = ['ID','ID PROPIETARIO','TIPO','DESCRIPCIÓN','DIRECCIÓN','SECTOR','CIUDAD/PROVINCIA',
      'LATITUD','LONGITUD','MONTO ALQUILER','ESTADO','NIVEL MORA','REFERENCIA INTERNA','FECHA REG']
write_hdr(ws1, 2, H1)
for i, r in enumerate(all_inmuebles, 3): write_row(ws1, i, r, i%2==0)
set_widths(ws1, 'ABCDEFGHIJKLMN', [16,16,12,42,28,14,22,12,12,16,12,12,22,12])
ws1.freeze_panes = 'A3'

# INQUILINOS
ws2 = wb.create_sheet('👥 INQUILINOS')
ws2.sheet_view.showGridLines = False
banner(ws2, f'INQUILINOS — {len(inquilinos)} personas registradas', 11)
H2 = ['ID','NOMBRE COMPLETO','CÉDULA','TELÉFONO','EMAIL','OCUPACIÓN',
      'REF PERSONAL','TEL REFERENCIA','ESTADO','NOTAS','FECHA REG']
write_hdr(ws2, 2, H2)
for i, r in enumerate(inquilinos, 3): write_row(ws2, i, r, i%2==0)
set_widths(ws2, 'ABCDEFGHIJK', [16,32,13,13,22,16,22,14,10,22,12])
ws2.freeze_panes = 'A3'

# CONTRATOS
ws3 = wb.create_sheet('📋 CONTRATOS')
ws3.sheet_view.showGridLines = False
banner(ws3, f'CONTRATOS — {len(contratos)} contratos vigentes (30 días gracia · mora 5%)', 16)
H3 = ['ID','ID INMUEBLE','ID INQUILINO','FECHA INICIO','FECHA FIN','MONTO',
      'PERIODICIDAD','DÍA COBRO','DÍAS GRACIA','TIPO MORA','MORA %','MORA FIJA',
      'DEPÓSITO','ESTADO','NOTAS','FECHA REG']
write_hdr(ws3, 2, H3)
for i, r in enumerate(contratos, 3): write_row(ws3, i, r, i%2==0)
set_widths(ws3, 'ABCDEFGHIJKLMNOP', [16,16,16,12,12,14,12,10,11,12,9,10,12,10,18,12])
ws3.freeze_panes = 'A3'

# CXC PENDIENTES
ws4 = wb.create_sheet('💰 CXC_PENDIENTES')
ws4.sheet_view.showGridLines = False
banner(ws4, f'CxC PENDIENTES — {len(cxc_pendientes)} cargos sin pagar al {HOY} (incluye Junio 2026)', 19, VERDE_H)
H4 = ['ID CXC','ID CONTRATO','ID INQUILINO','NOMBRE','ID INMUEBLE',
      'TIPO','DESCRIPCIÓN','PERÍODO INI','PERÍODO FIN','FECHA VENC',
      'MONTO ORIG','MORA','PAGADO','SALDO','ESTADO','DÍAS MORA',
      'ÚLT PAGO','NOTAS','FECHA REG']
write_hdr(ws4, 2, H4, bg=VERDE_H)
for i, r in enumerate(cxc_pendientes, 3): write_row(ws4, i, r, i%2==0)
set_widths(ws4, 'ABCDEFGHIJKLMNOPQRS', [16,16,16,26,16,12,26,12,12,12,14,12,12,14,11,10,12,18,12])
ws4.freeze_panes = 'A3'

# PAGOS HISTÓRICOS
ws5 = wb.create_sheet('✅ PAGOS_HISTORICOS')
ws5.sheet_view.showGridLines = False
banner(ws5, f'PAGOS — {len(cxc_pagadas)} pagos Ene-Jun 2026 · Ene-Mar=LIQUIDADO · Abr-Jun=EN_ADMINISTRACION', 17, PURP_H)
H5 = ['ID PAGO','ID CXC','ID CONTRATO','ID INQUILINO','NOMBRE','ID INMUEBLE',
      'SESIÓN','FECHA PAGO','MONTO ALQUILER','MORA','TOTAL PAGADO',
      'FORMA PAGO','REFERENCIA','ESTADO CUSTODIA','ID LIQUIDACIÓN','NOTAS','FECHA REG']
write_hdr(ws5, 2, H5, bg=PURP_H)
for i, r in enumerate(cxc_pagadas, 3): write_row(ws5, i, r, i%2==0)
set_widths(ws5, 'ABCDEFGHIJKLMNOPQ', [16,16,16,16,26,16,12,12,16,12,14,12,16,20,14,22,12])
ws5.freeze_panes = 'A3'

# RESUMEN
ws6 = wb.create_sheet('📊 RESUMEN')
ws6.sheet_view.showGridLines = False
ws6.column_dimensions['A'].width = 42; ws6.column_dimensions['B'].width = 20
ws6.merge_cells('A1:B1')
c = ws6['A1']
c.value = 'RESUMEN — Carga de Datos Reales de Prueba'
c.font = fn(bold=True, color='FFFFFF', size=12); c.fill = hf(AZUL_DARK); c.alignment = al_c()
ws6.row_dimensions[1].height = 30

filas_res = [
    ('Propietarios (13 pestañas del archivo)',       len(propietarios)),
    ('Inmuebles OCUPADOS',                           n_ocup),
    ('Inmuebles DISPONIBLES (unidades vacías)',       n_disp),
    ('Total inmuebles',                              len(all_inmuebles)),
    ('Inquilinos (personas que pagan)',               len(inquilinos)),
    ('Contratos vigentes 2026',                      len(contratos)),
    ('CxC Pendientes / En mora (incl. Junio)',        len(cxc_pendientes)),
    ('Pagos históricos registrados',                 len(cxc_pagadas)),
    ('', ''),
    ('CONDICIONES APLICADAS', ''),
    ('Días de gracia antes de mora', '30 días'),
    ('Porcentaje de mora', '5% del alquiler'),
    ('Día de corte (cobro mensual)', 'Día 1 de cada mes'),
    ('Forma de pago asumida', 'EFECTIVO'),
    ('', ''),
    ('CUSTODIA DE PAGOS', ''),
    ('Ene - Mar 2026',  'LIQUIDADO'),
    ('Abr - Jun 2026',  'EN_ADMINISTRACION'),
    ('', ''),
    ('PROVINCIAS GPS ASIGNADAS', ''),
    ('Propietarios 1,4,7,10,13', 'La Romana'),
    ('Propietarios 2,5,8,11',    'La Altagracia'),
    ('Propietarios 3,6,9,12',    'San Pedro de Macorís'),
    ('', ''),
    ('⚠ PARA ELIMINAR DATOS DE PRUEBA', ''),
    ('Filtrar en Google Sheets por columna ID que empiece con TEST-', ''),
    ('Seleccionar todas → Eliminar filas', ''),
]
for i, (a_val, b_val) in enumerate(filas_res, 2):
    a = ws6.cell(row=i, column=1, value=a_val)
    b = ws6.cell(row=i, column=2, value=b_val)
    is_hdr = str(b_val) == '' and a_val not in ('', '⚠ PARA ELIMINAR DATOS DE PRUEBA')
    a.font = fn(bold=is_hdr or 'PARA ELIMINAR' in str(a_val), size=10)
    b.font = fn(bold=True, color=AZUL_CORP, size=10)
    a.alignment = al_l(); b.alignment = al_c()
    if a_val and b_val != '': a.border = brd(); b.border = brd()
    if 'PARA ELIMINAR' in str(a_val):
        a.fill = hf('FFCDD2'); b.fill = hf('FFCDD2')
    if is_hdr:
        a.fill = hf(AZUL_CORP); a.font = fn(bold=True, color='FFFFFF', size=10)

out = r'C:\Personales RM\Claude\Sistema de Condominios\web-app\CARGA_REAL_PRUEBAS.xlsx'
wb.save(out)
print(f'Archivo: {out}')
print(f'  Propietarios   : {len(propietarios)}')
print(f'  Inmuebles ocup : {n_ocup}')
print(f'  Inmuebles disp : {n_disp}')
print(f'  Inquilinos     : {len(inquilinos)}')
print(f'  Contratos      : {len(contratos)}')
print(f'  CxC Pendientes : {len(cxc_pendientes)}')
print(f'  Pagos hist.    : {len(cxc_pagadas)}')
