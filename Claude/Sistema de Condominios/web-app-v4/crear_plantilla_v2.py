# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

OUT = r'C:\Users\Rudy\OneDrive - AM COMERCIAL\Rudy\Claude\Sistema de Condominios\web-app\Plantilla_Carga_Maestras_InmueblesRD_v2.xlsx'

# ── Paleta ──────────────────────────────────────────────────────────────────
AZUL       = '1B4F8C'
AZUL_FONT  = 'FFFFFF'
AMARILLO   = 'FFFACD'
GRIS_INS   = 'F1F5F9'
GRIS_ID    = 'E2E8F0'
VERDE_TAB  = '16A34A'
MORADO_TAB = '7C3AED'
NARANJA_TAB= 'EA580C'
BORDE_CLR  = 'CBD5E1'
INICIO_FILL= 'DCFCE7'   # verde clarito para marcar "aquí empiezas"

def thin_border():
    s = Side(style='thin', color=BORDE_CLR)
    return Border(left=s, right=s, top=s, bottom=s)

def hdr_font():  return Font(name='Arial', bold=True, color=AZUL_FONT, size=10)
def body_font(): return Font(name='Arial', size=10)
def ins_font():  return Font(name='Arial', italic=True, size=9, color='475569')

def hdr_fill():  return PatternFill('solid', fgColor=AZUL)
def ej_fill():   return PatternFill('solid', fgColor=AMARILLO)
def ins_fill():  return PatternFill('solid', fgColor=GRIS_INS)
def id_fill():   return PatternFill('solid', fgColor=GRIS_ID)

def center(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def left():   return Alignment(horizontal='left',   vertical='center', wrap_text=True)

# ── Helpers ──────────────────────────────────────────────────────────────────
def add_sheet(wb, name, tab_color):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab_color
    return ws

def write_instruction_row(ws, col_count, text):
    ws.row_dimensions[1].height = 46
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    c = ws.cell(1, 1, text)
    c.font  = ins_font()
    c.fill  = ins_fill()
    c.alignment = left()
    c.border = thin_border()

def write_header_row(ws, cols, row=2):
    ws.row_dimensions[row].height = 32
    for ci, (hdr, comment_txt) in enumerate(cols, 1):
        c = ws.cell(row, ci, hdr)
        c.font      = hdr_font()
        c.fill      = hdr_fill()
        c.alignment = center()
        c.border    = thin_border()
        if comment_txt:
            c.comment = Comment(comment_txt, 'Sistema')

def write_example_rows(ws, rows, start_row=3, id_col=1):
    for ri, row_data in enumerate(rows):
        r = start_row + ri
        ws.row_dimensions[r].height = 20
        for ci, val in enumerate(row_data, 1):
            c = ws.cell(r, ci, val)
            fill = id_fill() if ci == id_col else ej_fill()
            c.fill      = fill
            c.font      = body_font()
            c.alignment = left()
            c.border    = thin_border()
            if ci == id_col:
                c.font = Font(name='Arial', size=10, color='94A3B8', strike=True, italic=True)

def mark_inicio_real(ws, col_count, row=5):
    """Marca visualmente la fila donde debe empezar la digitación real."""
    ws.row_dimensions[row].height = 18
    for ci in range(1, col_count + 1):
        c = ws.cell(row, ci)
        c.fill = PatternFill('solid', fgColor=INICIO_FILL)
        c.border = Border(top=Side(style='medium', color='16A34A'),
                           left=Side(style='thin', color=BORDE_CLR),
                           right=Side(style='thin', color=BORDE_CLR),
                           bottom=Side(style='thin', color=BORDE_CLR))
    c = ws.cell(row, 1)
    c.value = '⬇ FILA 5: empieza a digitar tus datos reales aquí ⬇' if col_count > 1 else ''
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=min(col_count, 6))
    c.font = Font(name='Arial', italic=True, bold=True, size=8, color='15803D')
    c.alignment = left()

def set_col_widths(ws, widths):
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(ci)].width = w

def add_autofilter(ws, col_count, row=2):
    ws.auto_filter.ref = f"A{row}:{get_column_letter(col_count)}{row}"

def freeze_at(ws, row=5):
    ws.freeze_panes = f'A{row}'

def add_dv(ws, formula, col, first_row=5, last_row=2000):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True,
                        showErrorMessage=True, errorTitle='Valor inválido',
                        error='Seleccione un valor de la lista')
    ws.add_data_validation(dv)
    col_l = get_column_letter(col)
    dv.sqref = f'{col_l}{first_row}:{col_l}{last_row}'


# ════════════════════════════════════════════════════════════════════════════
wb = Workbook()
wb.remove(wb.active)

# ══════════════════════════════════════════════════════════════════════════
# HOJA 0 – INSTRUCCIONES
# ══════════════════════════════════════════════════════════════════════════
ws0 = add_sheet(wb, '📋 INSTRUCCIONES', VERDE_TAB)
ws0.column_dimensions['A'].width = 4
ws0.column_dimensions['B'].width = 6
ws0.column_dimensions['C'].width = 55
ws0.column_dimensions['D'].width = 35
ws0.column_dimensions['E'].width = 35

ws0.row_dimensions[1].height = 12
ws0.row_dimensions[2].height = 52
ws0.merge_cells('B2:E2')
c = ws0['B2']
c.value     = '📊  Plantilla de Carga de Maestras (v2)\nSistema de Administración de Inmuebles RD'
c.font      = Font(name='Arial', bold=True, size=18, color=AZUL)
c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

ws0.row_dimensions[3].height = 22
ws0.merge_cells('B3:E3')
c = ws0['B3']
c.value     = 'Incluye las mejoras: Co-Propietarios (Familia), Adelantos a propietarios y mora/gracia en Contratos.'
c.font      = Font(name='Arial', italic=True, size=11, color='475569')
c.alignment = Alignment(horizontal='left', vertical='center')

ws0.row_dimensions[4].height = 8

ws0.row_dimensions[5].height = 24
ws0.merge_cells('B5:E5')
c = ws0['B5']
c.value = 'INSTRUCCIONES DE USO'
c.font  = Font(name='Arial', bold=True, size=12, color=AZUL)
c.fill  = PatternFill('solid', fgColor='EFF6FF')
c.alignment = Alignment(horizontal='left', vertical='center')
c.border = thin_border()

pasos = [
    ('1', 'Orden de carga',
     'Llena las hojas en este orden: PROPIETARIOS → CO_PROPIETARIOS (si aplica) → INQUILINOS → INMUEBLES → '
     'CONTRATOS → CxC_CARGOS → ADELANTOS. Los IDs de una hoja son referenciados por las siguientes.'),
    ('2', '⚠️ LA CAPTURA REAL EMPIEZA EN LA FILA 5',
     'En cada hoja de datos: la fila 2 tiene los encabezados, las filas 3 y 4 son EJEMPLOS de referencia '
     '(con fondo amarillo) y NO se cargan al sistema. Debes escribir tus datos reales desde la FILA 5 en adelante '
     '(marcada con una línea verde). Todo lo que esté antes de la fila 5 se ignora en la importación.'),
    ('3', 'Campos obligatorios',
     'Los campos con (*) son obligatorios. Sin ellos el sistema rechazará el registro.'),
    ('4', 'Campo Id',
     'La columna "Id" se genera AUTOMÁTICAMENTE por el sistema. Déjala VACÍA. '
     'Aparece en gris/tachado como recordatorio visual.'),
    ('5', 'Formatos de fecha',
     'Las fechas deben estar en formato YYYY-MM-DD (ej: 2026-06-15). '
     'Formatea las celdas como Texto si Excel las convierte automáticamente.'),
    ('6', 'Listas desplegables',
     'Los campos con dropdown solo aceptan los valores de la lista. '
     'Haz clic en la celda para ver las opciones disponibles.'),
    ('7', 'Coordenadas GPS',
     'Para obtener latitud/longitud: abre Google Maps → clic derecho en la ubicación → '
     'copia el primer número (latitud) y el segundo (longitud, negativo para RD).'),
    ('8', 'Formato de cédula',
     'Ingresa las cédulas en formato 000-0000000-0 incluyendo los guiones. '
     'Formatea esas celdas como Texto para evitar que Excel elimine los ceros.'),
    ('9', 'IDs de referencia',
     'En las hojas que dependen de otra (CO_PROPIETARIOS, INMUEBLES, CONTRATOS, CxC_CARGOS, ADELANTOS) '
     'deberás colocar los IDs ya generados por el sistema (ej: PROP-001, INQ-001). Primero carga la maestra '
     'anterior y obtén sus IDs antes de continuar con la siguiente.'),
    ('10', 'Propietario tipo FAMILIAR',
     'En PROPIETARIOS, si TipoDueno = FAMILIAR, ese propietario debe tener miembros en la hoja '
     'CO_PROPIETARIOS. La suma de "PesoPorc" de todos sus miembros ACTIVOS debe ser exactamente 100%.'),
    ('11', 'Entrega',
     'Una vez completa, comparte este archivo con el administrador del sistema '
     'para proceder con la importación masiva de datos.'),
]

for i, (num, titulo, desc) in enumerate(pasos):
    r = 6 + i
    ws0.row_dimensions[r].height = 48
    c_num = ws0.cell(r, 2, num)
    c_num.font      = Font(name='Arial', bold=True, size=11, color=AZUL_FONT)
    c_num.fill      = hdr_fill()
    c_num.alignment = center()
    c_num.border    = thin_border()

    destacar = num == '2'
    c_tit = ws0.cell(r, 3, titulo)
    c_tit.font      = Font(name='Arial', bold=True, size=10, color='B91C1C' if destacar else AZUL)
    c_tit.fill      = PatternFill('solid', fgColor='FEF2F2' if destacar else 'EFF6FF')
    c_tit.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c_tit.border    = thin_border()

    ws0.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    c_desc = ws0.cell(r, 4, desc)
    c_desc.font      = Font(name='Arial', size=9, color='334155', bold=destacar)
    c_desc.fill      = PatternFill('solid', fgColor='FEF2F2' if destacar else 'F8FAFC')
    c_desc.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c_desc.border    = thin_border()

fila_tabla = 6 + len(pasos) + 1
ws0.row_dimensions[fila_tabla - 1].height = 12

ws0.row_dimensions[fila_tabla].height = 24
hdrs_res = ['Hoja', 'Descripción', 'Campos obligatorios']
for ci, h in enumerate(hdrs_res, 3):
    c = ws0.cell(fila_tabla, ci, h)
    c.font = hdr_font(); c.fill = hdr_fill()
    c.alignment = center(); c.border = thin_border()

resumen = [
    ('PROPIETARIOS',     'Dueños de los inmuebles, con datos bancarios, comisión y tipo de dueño (Único/Familiar)',
     'Nombre, TipoComision, TipoDueno'),
    ('CO_PROPIETARIOS',  'Miembros de la familia cuando el propietario es de tipo FAMILIAR, con su % de participación',
     'IdPropietario, Nombre, PesoPorc'),
    ('INQUILINOS',       'Personas que arriendan los inmuebles',
     'Nombre'),
    ('INMUEBLES',        'Propiedades disponibles para alquiler',
     'IdPropietario, Tipo, Descripcion, Direccion, Ciudad, MontoAlquiler'),
    ('CONTRATOS',        'Acuerdos de arrendamiento, incluyendo días de gracia y configuración de mora',
     'IdInmueble, IdInquilino, FechaInicio, FechaFin, MontoAlquiler, Periodicidad, DiaCobro, DiasGracia, TipoMora'),
    ('CxC_CARGOS',       'Cuentas por cobrar / cargos generados por contrato',
     'IdInquilino, IdInmueble, TipoCargo, Descripcion, PeriodoInicio, PeriodoFin, FechaVencimiento, MontoOriginal'),
    ('ADELANTOS',        'Pagos anticipados entregados a un propietario, descontados en su próxima liquidación',
     'IdPropietario, Fecha, Concepto, Monto'),
]

for ri, (h, d, obl) in enumerate(resumen):
    r = fila_tabla + 1 + ri
    ws0.row_dimensions[r].height = 34
    for ci, val in enumerate([h, d, obl], 3):
        c = ws0.cell(r, ci, val)
        c.font      = Font(name='Arial', size=9, color='1E293B', bold=(ci == 3))
        c.fill      = PatternFill('solid', fgColor='F8FAFC' if ri % 2 == 0 else 'FFFFFF')
        c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        c.border    = thin_border()

fila_warn = fila_tabla + 1 + len(resumen) + 1
ws0.row_dimensions[fila_warn - 1].height = 12
ws0.row_dimensions[fila_warn].height = 40
ws0.merge_cells(f'B{fila_warn}:E{fila_warn}')
c = ws0[f'B{fila_warn}']
c.value     = ('⚠️  IMPORTANTE: No cambies los nombres de las hojas ni el orden de las columnas. '
               'La captura real empieza en la FILA 5 de cada hoja — todo lo anterior (encabezados y ejemplos) no se importa.')
c.font      = Font(name='Arial', bold=True, size=10, color='92400E')
c.fill      = PatternFill('solid', fgColor='FEF3C7')
c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
c.border    = Border(
    left  =Side(style='medium', color='F59E0B'),
    right =Side(style='medium', color='F59E0B'),
    top   =Side(style='medium', color='F59E0B'),
    bottom=Side(style='medium', color='F59E0B'),
)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 1 – PROPIETARIOS
# ══════════════════════════════════════════════════════════════════════════
ws1 = add_sheet(wb, '👤 PROPIETARIOS', AZUL)

cols1 = [
    ('Id',                 'Auto-generado por el sistema. Dejar VACÍO.'),
    ('Nombre *',           'Nombre completo del propietario. Si es FAMILIAR, usa "Familia <Apellido>". Requerido.'),
    ('Cedula',             'Cédula dominicana. Formato: 000-0000000-0\nFormatea la celda como Texto.'),
    ('Telefono',           'Teléfono. Formato: 809-000-0000'),
    ('Email',              'Correo electrónico del propietario.'),
    ('Banco',              'Banco donde tiene la cuenta. Seleccionar de la lista.'),
    ('NumeroCuenta',       'Número de cuenta bancaria completo.'),
    ('TipoCuenta',         'Tipo de cuenta bancaria. Lista: AHORRO / CORRIENTE'),
    ('TipoComision',       'Cómo se calcula la comisión. Lista: PORCENTAJE / FIJO / MIXTO'),
    ('ComisionPorcentaje', 'Porcentaje de comisión (ej: 10 = 10%). Solo si TipoComision = PORCENTAJE o MIXTO. Poner 0 si no aplica.'),
    ('ComisionFija',       'Monto fijo de comisión en RD$. Solo si TipoComision = FIJO o MIXTO. Poner 0 si no aplica.'),
    ('TipoDueno *',        'NUEVO. Lista: ÚNICO / FAMILIAR.\nSi es FAMILIAR, registra a los miembros en la hoja CO_PROPIETARIOS '
                            'con su % de participación (deben sumar 100%).'),
    ('Notas',              'Observaciones adicionales del propietario. Campo libre.'),
]

write_instruction_row(ws1, len(cols1),
    'ℹ️  PROPIETARIOS — Registra los dueños de los inmuebles. Llena esta hoja PRIMERO. '
    'Los campos marcados con (*) son obligatorios. La columna "Id" se genera automáticamente — déjala vacía. '
    'Si TipoDueno = FAMILIAR, completa también la hoja CO_PROPIETARIOS. RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws1, cols1)

ej1 = [
    ['', 'José Mercedes Báez', '001-1234567-8', '809-555-1234', 'jose.mercedes@gmail.com',
     'BanReservas', '3100012345678', 'AHORRO', 'PORCENTAJE', 10, 0, 'ÚNICO',
     'Propietario del edificio Torre Azul, La Romana'],
    ['', 'Familia Rodríguez Pérez', '002-9876543-2', '809-777-5678', 'familia.rodriguez@hotmail.com',
     'Popular', '7200098765432', 'CORRIENTE', 'FIJO', 0, 5000, 'FAMILIAR',
     'Propiedad heredada, 3 hermanos co-propietarios (ver hoja CO_PROPIETARIOS)'],
]
write_example_rows(ws1, ej1)
mark_inicio_real(ws1, len(cols1))

set_col_widths(ws1, [12, 28, 18, 16, 30, 14, 18, 13, 15, 18, 15, 13, 40])
add_autofilter(ws1, len(cols1))
freeze_at(ws1, 5)

add_dv(ws1, '"BanReservas,Banesco,Popular,BHD León,Scotiabank,Banistmo,Asociación Cibao,Otra"', 6)
add_dv(ws1, '"AHORRO,CORRIENTE"', 8)
add_dv(ws1, '"PORCENTAJE,FIJO,MIXTO"', 9)
add_dv(ws1, '"ÚNICO,FAMILIAR"', 12)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 2 – CO_PROPIETARIOS (NUEVO)
# ══════════════════════════════════════════════════════════════════════════
ws2 = add_sheet(wb, '👨‍👩‍👧 CO_PROPIETARIOS', MORADO_TAB)

cols2 = [
    ('Id',            'Auto-generado por el sistema. Dejar VACÍO.'),
    ('IdPropietario *','ID del propietario FAMILIAR al que pertenece este miembro.\nEj: PROP-002\nDebe existir en la hoja PROPIETARIOS con TipoDueno = FAMILIAR.'),
    ('Nombre *',      'Nombre completo del miembro de la familia.'),
    ('Cedula',        'Cédula dominicana. Formato: 000-0000000-0'),
    ('Telefono',      'Teléfono del miembro.'),
    ('Email',         'Correo electrónico del miembro.'),
    ('Banco',         'Banco donde recibe sus liquidaciones (puede ser distinto al del propietario principal).'),
    ('NumeroCuenta',  'Número de cuenta bancaria del miembro.'),
    ('TipoCuenta',    'Tipo de cuenta. Lista: AHORRO / CORRIENTE'),
    ('PesoPorc *',    'Porcentaje de participación de este miembro (ej: 33.33).\nLa suma de todos los miembros ACTIVOS de un mismo propietario debe ser EXACTAMENTE 100.'),
    ('Notas',         'Observaciones. Ej: parentesco, acuerdos especiales.'),
]

write_instruction_row(ws2, len(cols2),
    'ℹ️  CO_PROPIETARIOS — Solo para propietarios con TipoDueno = FAMILIAR. Registra a cada miembro de la familia y su % '
    'de participación (PesoPorc). La suma de todos los miembros de un mismo propietario debe ser exactamente 100. '
    'RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws2, cols2)

ej2 = [
    ['', 'PROP-002', 'Carlos Rodríguez Pérez', '002-1111111-1', '809-111-1111',
     'carlos.rodriguez@gmail.com', 'Popular', '7200011111111', 'AHORRO', 40, 'Hermano mayor, administra la cuenta principal'],
    ['', 'PROP-002', 'Ana Rodríguez Pérez', '002-2222222-2', '809-222-2222',
     'ana.rodriguez@gmail.com', 'BHD León', '5300022222222', 'CORRIENTE', 30, 'Hermana, vive en el extranjero'],
]
write_example_rows(ws2, ej2)
mark_inicio_real(ws2, len(cols2))

set_col_widths(ws2, [12, 16, 26, 18, 16, 28, 14, 18, 13, 12, 38])
add_autofilter(ws2, len(cols2))
freeze_at(ws2, 5)

add_dv(ws2, '"AHORRO,CORRIENTE"', 9)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 3 – INQUILINOS
# ══════════════════════════════════════════════════════════════════════════
ws3 = add_sheet(wb, '👥 INQUILINOS', AZUL)

cols3 = [
    ('Id',                  'Auto-generado por el sistema. Dejar VACÍO.'),
    ('Nombre *',            'Nombre completo del inquilino. Requerido.'),
    ('Cedula',              'Cédula dominicana. Formato: 000-0000000-0\nFormatea la celda como Texto.'),
    ('Telefono',            'Teléfono principal del inquilino.'),
    ('Email',               'Correo electrónico del inquilino.'),
    ('Ocupacion',           'Ocupación o empresa donde trabaja el inquilino.'),
    ('ReferenciaPersonal',  'Nombre completo de una persona de referencia.'),
    ('TelefonoReferencia',  'Teléfono de la persona de referencia.'),
    ('Notas',               'Observaciones adicionales. Campo libre.'),
]

write_instruction_row(ws3, len(cols3),
    'ℹ️  INQUILINOS — Registra las personas que arriendan los inmuebles. '
    'Llena esta hoja DESPUÉS de PROPIETARIOS. La columna "Id" se genera automáticamente — déjala vacía. '
    'RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws3, cols3)

ej3 = [
    ['', 'Ana Yolanda Fernández Cruz', '402-1234567-8', '809-444-2345',
     'ana.fernandez@gmail.com', 'Maestra del MINERD', 'Pedro García', '829-111-2222',
     'Inquilina puntual, lleva 3 años en el apartamento'],
    ['', 'Luis Antonio Tejada Méndez', '222-9876543-2', '829-666-7890',
     'luistejada@hotmail.com', 'Empleado de Claro RD', 'Carmen Méndez', '809-333-4444',
     'Contrato renovado por segundo año consecutivo'],
]
write_example_rows(ws3, ej3)
mark_inicio_real(ws3, len(cols3))

set_col_widths(ws3, [12, 28, 18, 16, 30, 30, 25, 18, 40])
add_autofilter(ws3, len(cols3))
freeze_at(ws3, 5)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 4 – INMUEBLES
# ══════════════════════════════════════════════════════════════════════════
ws4 = add_sheet(wb, '🏢 INMUEBLES', AZUL)

cols4 = [
    ('Id',            'Auto-generado por el sistema. Dejar VACÍO.'),
    ('IdPropietario *', 'ID del propietario dueño del inmueble.\nEj: PROP-001\nDebe existir en la hoja PROPIETARIOS o en el sistema.'),
    ('Tipo *',        'Tipo de inmueble. Seleccionar de la lista:\nAPARTAMENTO / CASA / LOCAL / OFICINA / VILLA / HABITACION / OTRO'),
    ('Descripcion *', 'Descripción clara del inmueble.\nEj: Apto 2BR Torre Caribeña Piso 3'),
    ('Direccion *',   'Dirección física completa del inmueble.'),
    ('Sector',        'Sector, barrio o urbanización donde está ubicado.'),
    ('Ciudad *',      'Ciudad. Seleccionar de la lista.'),
    ('Latitud',       'Coordenada GPS de latitud.\nEj: 18.427000\nPara RD los valores están entre 17.5 y 19.9\nSi no la tienes, en el sistema puedes elegirla en el mapa al crear el inmueble.'),
    ('Longitud',      'Coordenada GPS de longitud.\nEj: -68.972000\nPara RD los valores son NEGATIVOS entre -68 y -72'),
    ('MontoAlquiler *', 'Monto mensual del alquiler en pesos dominicanos (RD$). Sin símbolos ni comas.'),
    ('Notas',         'Observaciones adicionales del inmueble. Campo libre.'),
]

write_instruction_row(ws4, len(cols4),
    'ℹ️  INMUEBLES — Registra cada propiedad disponible para alquiler. '
    'Llena esta hoja DESPUÉS de PROPIETARIOS. Para las coordenadas GPS: en Google Maps, clic derecho en la ubicación y copia los números '
    '(o usa el selector de mapa dentro del sistema). RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws4, cols4)

ej4 = [
    ['', 'PROP-001', 'APARTAMENTO', 'Apto 2BR Torre Caribeña Piso 3',
     'C/ Duarte #45 Edif. Torre Caribeña', 'Centro', 'La Romana',
     18.427000, -68.972000, 25000, 'Incluye agua y electricidad básica hasta RD$800'],
    ['', 'PROP-002', 'LOCAL', 'Local Comercial Planta Baja Av. Libertad',
     'Av. Libertad #78', 'El Seibo', 'La Altagracia',
     18.765000, -68.723000, 45000, 'Zona comercial de alta circulación. 80 m².'],
]
write_example_rows(ws4, ej4)
mark_inicio_real(ws4, len(cols4))

set_col_widths(ws4, [12, 15, 14, 32, 35, 20, 16, 12, 12, 15, 38])
add_autofilter(ws4, len(cols4))
freeze_at(ws4, 5)

add_dv(ws4, '"APARTAMENTO,CASA,LOCAL,OFICINA,VILLA,HABITACION,OTRO"', 3)
add_dv(ws4, '"La Romana,La Altagracia,San Pedro de Macorís,Santo Domingo,Santiago,Puerto Plata,Otra"', 7)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 5 – CONTRATOS
# ══════════════════════════════════════════════════════════════════════════
ws5 = add_sheet(wb, '📋 CONTRATOS', AZUL)

cols5 = [
    ('Id',             'Auto-generado por el sistema. Dejar VACÍO.'),
    ('IdInmueble *',   'ID del inmueble arrendado.\nEj: INM-001\nDebe existir en la hoja INMUEBLES o en el sistema.'),
    ('IdInquilino *',  'ID del inquilino que firma el contrato.\nEj: INQ-001\nDebe existir en la hoja INQUILINOS o en el sistema.'),
    ('FechaInicio *',  'Fecha de inicio del contrato.\nFormato: YYYY-MM-DD\nEj: 2026-01-01'),
    ('FechaFin *',     'Fecha de vencimiento del contrato.\nFormato: YYYY-MM-DD\nEj: 2027-01-01'),
    ('MontoAlquiler *','Monto pactado de alquiler en RD$. Sin símbolos ni comas.'),
    ('Periodicidad *', 'Frecuencia de cobro. Lista: MENSUAL / QUINCENAL'),
    ('DiaCobro *',     'Día del mes en que vence el pago (si es MENSUAL).\nEj: 1 = día 1 de cada mes\nNúmero del 1 al 28.'),
    ('DiasGracia *',   'Días de gracia antes de aplicar mora.\nEj: 5 = 5 días después del vencimiento. Estándar del sistema: 5.'),
    ('TipoMora *',     'Cómo se calcula la mora. Lista: PORCENTAJE / MONTO_FIJO'),
    ('MoraPorcentaje', 'Porcentaje de mora POR PERÍODO atrasado (ej: 5 = 5%).\nSolo si TipoMora = PORCENTAJE. Poner 0 si no aplica.\n'
                        'El sistema la recalcula cada noche multiplicando por la cantidad de períodos vencidos.'),
    ('MontoMoraFija',  'Monto fijo de mora POR PERÍODO atrasado en RD$.\nSolo si TipoMora = MONTO_FIJO. Poner 0 si no aplica.'),
    ('Deposito',       'Monto del depósito de garantía cobrado al inquilino en RD$.\nGeneralmente equivale a 1 mes de alquiler.'),
    ('Notas',          'Observaciones adicionales del contrato. Campo libre.'),
]

write_instruction_row(ws5, len(cols5),
    'ℹ️  CONTRATOS — Registra los acuerdos de arrendamiento. Llena esta hoja DESPUÉS de INMUEBLES e INQUILINOS. '
    'Usa el formato YYYY-MM-DD para todas las fechas. El sistema genera automáticamente la próxima CxC de alquiler cada noche '
    'según Periodicidad y DiaCobro, y recalcula la mora según DiasGracia/TipoMora/MoraPorcentaje. '
    'RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws5, cols5)

ej5 = [
    ['', 'INM-001', 'INQ-001', '2026-01-01', '2027-01-01',
     25000, 'MENSUAL', 1, 5, 'PORCENTAJE', 5, 0, 25000,
     'Contrato anual renovado. Primera renovación.'],
    ['', 'INM-002', 'INQ-002', '2026-02-01', '2027-02-01',
     35000, 'QUINCENAL', 15, 3, 'MONTO_FIJO', 0, 1500, 35000,
     'Pago quincenal acordado por el inquilino. Mora fija RD$1,500 por quincena atrasada.'],
]
write_example_rows(ws5, ej5)
mark_inicio_real(ws5, len(cols5))

set_col_widths(ws5, [12, 14, 14, 14, 14, 15, 13, 10, 11, 13, 15, 14, 12, 38])
add_autofilter(ws5, len(cols5))
freeze_at(ws5, 5)

add_dv(ws5, '"MENSUAL,QUINCENAL"', 7)
add_dv(ws5, '"PORCENTAJE,MONTO_FIJO"', 10)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 6 – CxC_CARGOS
# ══════════════════════════════════════════════════════════════════════════
ws6 = add_sheet(wb, '💰 CxC_CARGOS', AZUL)

cols6 = [
    ('Id',               'Auto-generado por el sistema. Dejar VACÍO.'),
    ('IdContrato',       'ID del contrato asociado.\nEj: CTR-001\nPuede dejarse vacío si es un cargo manual sin contrato.'),
    ('IdInquilino *',    'ID del inquilino al que se le cobra.\nEj: INQ-001'),
    ('IdInmueble *',     'ID del inmueble al que aplica el cargo.\nEj: INM-001'),
    ('TipoCargo *',      'Tipo de cargo. Lista: ALQUILER / MORA / DEPOSITO / CARGO_EXTRA / OTRO'),
    ('Descripcion *',    'Descripción del cargo.\nEj: Alquiler enero 2026\nEj: Mora por retraso febrero 2026'),
    ('PeriodoInicio *',  'Inicio del período facturado.\nFormato: YYYY-MM-DD\nEj: 2026-01-01'),
    ('PeriodoFin *',     'Fin del período facturado.\nFormato: YYYY-MM-DD\nEj: 2026-01-31'),
    ('FechaVencimiento *','Fecha límite para pagar este cargo.\nFormato: YYYY-MM-DD'),
    ('MontoOriginal *',  'Monto total a cobrar en RD$. Sin símbolos ni comas.'),
    ('Estado *',         'Estado del cargo. Lista: PENDIENTE / EN_MORA / PAGADO / PAGADO_PARCIAL'),
    ('Notas',            'Observaciones adicionales. Campo libre.'),
]

write_instruction_row(ws6, len(cols6),
    'ℹ️  CxC_CARGOS — Cuentas por cobrar y cargos pendientes. Llena esta hoja DESPUÉS de CONTRATOS. '
    'Para alquileres regulares el sistema los genera automáticamente cada noche; usa esta hoja solo para cargos históricos o manuales. '
    'RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws6, cols6)

ej6 = [
    ['', 'CTR-001', 'INQ-001', 'INM-001', 'ALQUILER',
     'Alquiler enero 2026', '2026-01-01', '2026-01-31', '2026-01-01',
     25000, 'PENDIENTE', 'Pendiente de cobro mes de enero'],
    ['', '', 'INQ-002', 'INM-002', 'CARGO_EXTRA',
     'Pintura exterior solicitada por propietario', '2026-03-01', '2026-03-31', '2026-03-15',
     8000, 'PENDIENTE', 'Cargo manual autorizado por propietario el 2026-02-20'],
]
write_example_rows(ws6, ej6)
mark_inicio_real(ws6, len(cols6))

set_col_widths(ws6, [12, 12, 12, 12, 14, 38, 14, 14, 16, 15, 16, 38])
add_autofilter(ws6, len(cols6))
freeze_at(ws6, 5)

add_dv(ws6, '"ALQUILER,MORA,DEPOSITO,CARGO_EXTRA,OTRO"', 5)
add_dv(ws6, '"PENDIENTE,EN_MORA,PAGADO,PAGADO_PARCIAL"', 11)


# ══════════════════════════════════════════════════════════════════════════
# HOJA 7 – ADELANTOS (NUEVO)
# ══════════════════════════════════════════════════════════════════════════
ws7 = add_sheet(wb, '💸 ADELANTOS', NARANJA_TAB)

cols7 = [
    ('Id',             'Auto-generado por el sistema. Dejar VACÍO.'),
    ('IdPropietario *','ID del propietario que recibe el adelanto.\nEj: PROP-001\nDebe existir en la hoja PROPIETARIOS.'),
    ('Fecha *',        'Fecha en que se entregó el adelanto.\nFormato: YYYY-MM-DD'),
    ('Concepto *',     'Motivo del adelanto.\nEj: Adelanto para reparación de techo, pintura, etc.'),
    ('Monto *',        'Monto entregado en RD$. Sin símbolos ni comas.'),
    ('Notas',          'Observaciones adicionales. Campo libre.'),
]

write_instruction_row(ws7, len(cols7),
    'ℹ️  ADELANTOS — Pagos anticipados entregados a un propietario (para gastos o servicios) que se descuentan '
    'automáticamente en su próxima liquidación. Quedan en estado PENDIENTE hasta que se procese esa liquidación. '
    'RECUERDA: la captura real empieza en la FILA 5.')
write_header_row(ws7, cols7)

ej7 = [
    ['', 'PROP-001', '2026-05-10', 'Adelanto para reparación de techo', 15000,
     'Acordado verbalmente con el propietario, se descuenta en junio'],
    ['', 'PROP-002', '2026-05-20', 'Adelanto para pintura exterior del inmueble INM-002', 8000,
     'Solicitado por el propietario antes de la próxima liquidación'],
]
write_example_rows(ws7, ej7)
mark_inicio_real(ws7, len(cols7))

set_col_widths(ws7, [12, 16, 14, 42, 14, 40])
add_autofilter(ws7, len(cols7))
freeze_at(ws7, 5)


# ══════════════════════════════════════════════════════════════════════════
# Guardar
# ══════════════════════════════════════════════════════════════════════════
wb.active = ws0
wb.save(OUT)
print(f'OK: {OUT}')
